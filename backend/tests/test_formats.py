"""
test_formats.py - Ingesting files other than PDF: CSV, Excel, PowerPoint, Word, text, Markdown.
Offline: sample files are built in memory and the vision model is the fake from fakes.py.

Run (from backend/):   python -m pytest
"""

# --- Imports -----------------------------------------------------------------
import csv
import datetime
import io

import openpyxl
import pymupdf
import pytest
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches

from ingest.formats import SUPPORTED, check_format, read_document


# --- Sample files, built in memory -------------------------------------------------
def make_csv(rows: int = 3) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Particulars", "Q3 FY25", "Q2 FY25"])
    for i in range(rows):
        writer.writerow([f"Line item {i}", f"{1234.5 + i:,.1f}", "1,100.0"])
    return buffer.getvalue().encode()


def make_xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Income Statements"
    sheet.append(["(In millions, except per share amounts)", None, None])
    sheet.append(["Item", "2025", "2024"])
    sheet.append(["Revenue", 57167, 50000])
    sheet.append(["Net income", 17210.5, None])
    sheet.append(["Date", datetime.datetime(2025, 9, 30), "a | pipe"])
    other = workbook.create_sheet("Balance Sheets")
    other.append(["Assets", "2025"])
    other.append(["Cash", 6800])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx(tag: str = "deck") -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])          # title only
    slide.shapes.title.text = f"Financial summary ({tag})"   # the tag makes each deck a new file (new doc_id)
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "Revenue increased 18% year over year, driven by cloud growth."
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3.5), Inches(5), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "FY26 Q1"
    table.cell(1, 0).text, table.cell(1, 1).text = "Revenue", "$77.7"
    picture_only = deck.slides.add_slide(deck.slide_layouts[6])   # blank: a picture and nothing else
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 320, 180), False)
    pixmap.clear_with(200)
    picture_only.shapes.add_picture(io.BytesIO(pixmap.tobytes("png")), Inches(0), Inches(0), Inches(8), Inches(4.5))
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = WordDocument()
    document.add_heading("Board decision", level=1)
    document.add_paragraph("The board approved a dividend of 5 per share for the year.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "Dividend", "5.00"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def stored_text(vectors) -> str:
    return "\n".join(point.payload["text"] for point in vectors.points.values())


# =====================================================================================
# 1. Each format is read and stored
# =====================================================================================
def test_csv_rows_become_fact_cards_with_their_values(ingest, vectors):
    status, body = ingest(make_csv(), name="results.csv")
    assert status == 200 and body["job"]["status"] == "complete", body["job"].get("error")
    text = stored_text(vectors)
    assert "Line item 0" in text and "1,234.5" in text and "Q3 FY25" in text    # header is kept with the values


def test_xlsx_sheets_become_sections_and_cells_are_exact(ingest, vectors):
    status, body = ingest(make_xlsx(), name="statements.xlsx")
    assert status == 200 and body["job"]["status"] == "complete", body["job"].get("error")
    sections = {point.payload["section"] for point in vectors.points.values()}
    assert any("Income Statements" in s for s in sections) and any("Balance Sheets" in s for s in sections)
    text = stored_text(vectors)
    assert "57,167" in text and "17,210.5" in text and "6,800" in text             # stored values, separators kept
    assert "a \\| pipe" in text                                                    # a pipe inside a cell is escaped
    assert "2025-09-30" in text                                                    # dates as ISO


def test_pptx_text_slide_is_read_without_the_vision_model(ingest, vectors, llm):
    status, body = ingest(make_pptx(), name="slides.pptx")
    assert status == 200 and body["job"]["status"] == "complete", body["job"].get("error")
    text = stored_text(vectors)
    assert "Revenue increased 18%" in text and "$77.7" in text                    # text box and table
    assert "Financial summary" in text
    assert llm["vision"].calls["describe_image"] == 1                              # only the picture-only slide


def test_pptx_picture_slide_is_read_by_vision_once_then_cached(ingest, llm):
    ingest(make_pptx("cache-test"), name="slides.pptx")
    assert llm["vision"].calls["describe_image"] == 1
    status, body = ingest(make_pptx("cache-test"), name="slides.pptx", force=True)  # same file again
    assert status == 200 and body["job"]["status"] == "complete"
    assert llm["vision"].calls["describe_image"] == 1                              # cache: no second call


def test_docx_headings_paragraphs_and_tables(ingest, vectors):
    status, body = ingest(make_docx(), name="board.docx")
    assert status == 200 and body["job"]["status"] == "complete", body["job"].get("error")
    text = stored_text(vectors)
    assert "The board approved a dividend of 5 per share" in text
    assert "Dividend" in text and "5.00" in text
    assert any(point.payload["section"].startswith("Board decision") for point in vectors.points.values())


@pytest.mark.parametrize("name, data", [("notes.txt", b"First paragraph about results.\n\nSecond paragraph."),
                                        ("notes.md", b"## Outlook\n\nDemand stayed strong.")])
def test_text_and_markdown(ingest, vectors, name, data):
    status, body = ingest(data, name=name)
    assert status == 200 and body["job"]["status"] == "complete", body["job"].get("error")
    assert "Demand stayed strong" in stored_text(vectors) or "First paragraph" in stored_text(vectors)


# =====================================================================================
# 2. Long files are split into pages that keep their header
# =====================================================================================
def test_long_csv_is_split_into_pages_each_with_the_header():
    pages = read_document(make_csv(rows=130), "big.csv")
    assert len(pages) == 3                                                        # 60 + 60 + 10 rows
    assert all("Particulars" in page.text for page in pages)                      # header on every page
    assert [page.page_no for page in pages] == [1, 2, 3]


def test_long_text_is_split_by_size():
    text = ("Sentence number. " * 40 + "\n\n") * 10
    assert len(read_document(text.encode(), "long.txt")) > 1


# =====================================================================================
# 3. Wrong or unsupported files are refused before any work
# =====================================================================================
@pytest.mark.parametrize("name, data", [
    ("report.exe", b"MZ..."),                       # unsupported extension
    ("report.xlsx", make_csv()),                    # content does not match the extension
    ("report.csv", b"\xff\xfe\x00bad"),             # not UTF-8 text
    ("report.pdf", make_csv()),                     # a PDF must start with %PDF
])
def test_refuses_files_that_are_not_what_they_claim(client, name, data):
    reply = client.post("/ingest", files={"file": (name, data, "application/octet-stream")},
                        data={"company": "VIDEOIND", "period": "Q3 FY25"})
    assert reply.status_code == 400, reply.text


def test_factory_knows_every_format():
    assert set(SUPPORTED) == {".pdf", ".csv", ".xlsx", ".pptx", ".docx", ".txt", ".md"}
    from ingest.formats import get_reader
    assert get_reader("x.xlsx").name == "xlsx" and get_reader("x.MD").name == "text"
    with pytest.raises(ValueError, match="not supported"):
        get_reader("x.exe")


def test_every_supported_extension_is_accepted_by_the_checker():
    assert check_format("a.pdf", b"%PDF-1.4") == ".pdf"
    assert check_format("a.CSV", make_csv()) == ".csv"
    assert check_format("a.xlsx", make_xlsx()) == ".xlsx"
    assert check_format("a.pptx", make_pptx()) == ".pptx"
    assert check_format("a.docx", make_docx()) == ".docx"
    assert check_format("a.txt", b"hello") == ".txt"
    assert check_format("a.md", b"# hello") == ".md"


# =====================================================================================
# 4. Citations: spreadsheets and CSV show their sheet and rows, not a page number
# =====================================================================================
def test_spreadsheet_citation_shows_sheet_and_rows_not_a_page(client, llm, ingest):
    ingest(make_xlsx(), name="statements.xlsx")
    result = client.post("/ask", json={"question": "What was revenue?", "company": "VIDEOIND"}).json()
    assert result["sources"], "the fake answer cites S1, so a source is returned"
    source = result["sources"][0]
    assert source["location"].startswith("sheet '") and "rows " in source["location"]
    assert source["citation"] == source["location"] and source["pages"] == ""
    [prompt] = llm["answer"].prompts
    assert "statements.xlsx, sheet '" in prompt and "statements.xlsx, page" not in prompt


def test_csv_citation_shows_rows(client, llm, ingest):
    ingest(make_csv(rows=3), name="results.csv")
    result = client.post("/ask", json={"question": "What was revenue?", "company": "VIDEOIND"}).json()
    source = result["sources"][0]
    assert source["location"].startswith("rows ") and source["pages"] == ""


def test_pdf_and_slide_citations_still_show_the_page(client, llm, ingest):
    ingest(make_pptx("cite-test"), name="slides.pptx")
    result = client.post("/ask", json={"question": "What was revenue?", "company": "VIDEOIND"}).json()
    source = result["sources"][0]
    assert source["location"] == "" and source["citation"].startswith("page ")
