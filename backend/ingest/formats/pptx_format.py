"""
pptx_format.py - PowerPoint decks: one page per slide. Text, tables and speaker notes are
read directly. A slide with little text that holds a picture goes to the vision model
(the picture, cached like a scanned PDF page).
"""

# --- Imports -----------------------------------------------------------------
# io:            read the uploaded bytes.
# pymupdf:       turns a slide's picture into PNG for the vision model.
# python-pptx:   reads slides.
# MIN_TEXT_CHARS, PageResult: thresholds and page record (ingest/pdf_reader.py).
# get_llm, LLMBlocked: the vision model and its "gave no text" error (utils/llm).
# page_cache, prompts: cache for picture slides; the transcription prompt.
import io

from ingest.formats.base import DocumentReader
from ingest.formats.common import cell_text, check_office_zip, markdown_table
from ingest.pdf_reader import MIN_TEXT_CHARS, PageResult
from utils import page_cache, prompts
from utils.llm import LLMBlocked, get_llm

# Cache folder name for slide pictures (same settings -> same folder).
CACHE_VARIANT = page_cache.make_variant("pptx-picture", "vision", prompts.TRANSCRIBE_PAGE)


class PptxReader(DocumentReader):
    name = "pptx"
    extensions = (".pptx",)

    def check(self, data: bytes) -> None:
        check_office_zip(data, ".pptx")

    def read(self, data: bytes, doc_id: str, use_cache: bool):
        from pptx import Presentation  # lazy: only the Lambda-excluded ingest path needs it
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        deck = Presentation(io.BytesIO(data))
        pages = []
        for number, slide in enumerate(deck.slides, start=1):
            parts, pictures = [], []
            # Shapes top to bottom, then left to right, so text reads in order.
            for shape in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pictures.append(shape)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    parts.append(markdown_table([[cell_text(c.text) for c in row.cells] for row in shape.table.rows]))
                elif shape.has_text_frame:
                    text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                    if text:
                        parts.append(text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append("Speaker notes: " + notes)

            text = "\n\n".join(parts).strip()
            if pictures and len(text) < MIN_TEXT_CHARS:
                pages.append(self._read_picture(pictures[0], doc_id, number, use_cache, prefix=text))
            else:
                pages.append(PageResult(number, text, "text"))
        return pages

    def _read_picture(self, shape, doc_id: str, number: int, use_cache: bool, prefix: str) -> PageResult:
        """A picture slide -> vision model, or the cache. A failure is kept, never hidden."""
        if use_cache:
            cached = page_cache.load(doc_id, CACHE_VARIANT, number)
            if cached is not None:
                print(f"slide {number}: picture -> cache")
                return PageResult(number, _join(prefix, cached), "vision")
        print(f"slide {number}: picture -> vision model")
        import pymupdf  # lazy: only slides with pictures need it
        try:
            png = pymupdf.Pixmap(shape.image.blob).tobytes("png")
            markdown = get_llm("vision").describe_image(png, prompts.TRANSCRIBE_PAGE)
        except LLMBlocked as blocked:
            print(f"slide {number}:   FAILED: {blocked}")
            return PageResult(number, prefix, "failed", error=str(blocked))
        page_cache.save(doc_id, CACHE_VARIANT, number, markdown)
        return PageResult(number, _join(prefix, markdown), "vision")


def _join(prefix: str, markdown: str) -> str:
    return f"{prefix}\n\n{markdown}".strip() if prefix else markdown
